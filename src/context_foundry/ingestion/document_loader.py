"""
Document Loader for Context Foundry MVP2.
Loads documents from various formats: PDF, DOCX, MD, TXT.
"""
import os
import re
from pathlib import Path
from typing import Dict, List, Optional, Union
from dataclasses import dataclass
from datetime import datetime
import hashlib


@dataclass
class LoadedDocument:
    """Represents a loaded document with metadata."""
    id: str
    title: str
    content: str
    source_path: str
    file_type: str
    metadata: Dict
    loaded_at: datetime
    content_hash: str
    
    def to_dict(self) -> Dict:
        return {
            "id": self.id,
            "title": self.title,
            "content": self.content,
            "source_path": self.source_path,
            "file_type": self.file_type,
            "metadata": self.metadata,
            "loaded_at": self.loaded_at.isoformat(),
            "content_hash": self.content_hash,
        }


class DocumentLoader:
    """
    Loads documents from various file formats.
    Supports: PDF, DOCX, MD, TXT
    """
    
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt", ".markdown"}
    
    def __init__(self):
        self._pdf_available = self._check_pdf_support()
        self._docx_available = self._check_docx_support()
    
    def _check_pdf_support(self) -> bool:
        """Check if PDF parsing is available."""
        try:
            import pypdf
            return True
        except ImportError:
            return False
    
    def _check_docx_support(self) -> bool:
        """Check if DOCX parsing is available."""
        try:
            import docx
            return True
        except ImportError:
            return False
    
    def _compute_hash(self, content: str) -> str:
        """Compute MD5 hash of content for deduplication."""
        return hashlib.md5(content.encode()).hexdigest()
    
    def _generate_id(self, path: str, content_hash: str) -> str:
        """Generate deterministic document ID."""
        combined = f"{path}:{content_hash}"
        return hashlib.md5(combined.encode()).hexdigest()
    
    def _extract_title(self, content: str, filename: str) -> str:
        """Extract title from content or use filename."""
        lines = content.strip().split('\n')
        for line in lines[:5]:
            line = line.strip()
            if line.startswith('# '):
                return line[2:].strip()
            if line and not line.startswith('#') and len(line) > 10:
                return line[:100] if len(line) > 100 else line
        return Path(filename).stem.replace('_', ' ').replace('-', ' ').title()
    
    def load_text(self, path: str) -> LoadedDocument:
        """Load a plain text or markdown file."""
        with open(path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()
        
        file_type = "markdown" if path.endswith(('.md', '.markdown')) else "text"
        content_hash = self._compute_hash(content)
        
        return LoadedDocument(
            id=self._generate_id(path, content_hash),
            title=self._extract_title(content, path),
            content=content,
            source_path=path,
            file_type=file_type,
            metadata={
                "file_size": os.path.getsize(path),
                "line_count": content.count('\n') + 1,
                "word_count": len(content.split()),
            },
            loaded_at=datetime.now(),
            content_hash=content_hash,
        )
    
    def load_pdf(self, path: str) -> LoadedDocument:
        """Load a PDF file."""
        if not self._pdf_available:
            raise ImportError("pypdf is required for PDF support. Install with: pip install pypdf")
        
        import pypdf
        
        text_parts = []
        with open(path, 'rb') as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        
        content = '\n\n'.join(text_parts)
        content_hash = self._compute_hash(content)
        
        return LoadedDocument(
            id=self._generate_id(path, content_hash),
            title=self._extract_title(content, path),
            content=content,
            source_path=path,
            file_type="pdf",
            metadata={
                "file_size": os.path.getsize(path),
                "page_count": len(text_parts),
                "word_count": len(content.split()),
            },
            loaded_at=datetime.now(),
            content_hash=content_hash,
        )
    
    def load_docx(self, path: str) -> LoadedDocument:
        """Load a DOCX file."""
        if not self._docx_available:
            raise ImportError("python-docx is required for DOCX support. Install with: pip install python-docx")
        
        import docx
        
        doc = docx.Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = '\n\n'.join(paragraphs)
        content_hash = self._compute_hash(content)
        
        return LoadedDocument(
            id=self._generate_id(path, content_hash),
            title=self._extract_title(content, path),
            content=content,
            source_path=path,
            file_type="docx",
            metadata={
                "file_size": os.path.getsize(path),
                "paragraph_count": len(paragraphs),
                "word_count": len(content.split()),
            },
            loaded_at=datetime.now(),
            content_hash=content_hash,
        )
    
    def load(self, path: str) -> LoadedDocument:
        """Load a document from the given path."""
        path = str(path)
        
        if not os.path.exists(path):
            raise FileNotFoundError(f"Document not found: {path}")
        
        ext = Path(path).suffix.lower()
        
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {ext}. Supported: {self.SUPPORTED_EXTENSIONS}")
        
        if ext == ".pdf":
            return self.load_pdf(path)
        elif ext == ".docx":
            return self.load_docx(path)
        else:
            return self.load_text(path)
    
    def load_directory(
        self, 
        directory: str, 
        recursive: bool = True,
        extensions: Optional[List[str]] = None
    ) -> List[LoadedDocument]:
        """Load all supported documents from a directory."""
        directory = Path(directory)
        
        if not directory.exists():
            raise FileNotFoundError(f"Directory not found: {directory}")
        
        allowed_extensions = set(extensions) if extensions else self.SUPPORTED_EXTENSIONS
        
        documents = []
        pattern = "**/*" if recursive else "*"
        
        for path in directory.glob(pattern):
            if path.is_file() and path.suffix.lower() in allowed_extensions:
                try:
                    doc = self.load(str(path))
                    documents.append(doc)
                except Exception as e:
                    print(f"Warning: Failed to load {path}: {e}")
        
        return documents
    
    def load_from_string(
        self, 
        content: str, 
        title: str,
        source: str = "inline",
        file_type: str = "text"
    ) -> LoadedDocument:
        """Create a document from a string (useful for testing or API input)."""
        content_hash = self._compute_hash(content)
        
        return LoadedDocument(
            id=self._generate_id(source, content_hash),
            title=title,
            content=content,
            source_path=source,
            file_type=file_type,
            metadata={
                "word_count": len(content.split()),
                "line_count": content.count('\n') + 1,
            },
            loaded_at=datetime.now(),
            content_hash=content_hash,
        )
